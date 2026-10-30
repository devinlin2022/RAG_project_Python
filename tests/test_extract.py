"""Multi-format extraction: one page contract, six file types.

The samples are generated here instead of committed, so the repository stays
free of binary fixtures and the test proves the parsers, not the fixture.
"""

import pytest

from rag.ingest.extract import SUPPORTED_EXT, clean_text, extract_pages, format_label, supported


@pytest.fixture
def workdir(tmp_path):
    return tmp_path


def _make_docx(path):
    pytest.importorskip("docx")
    from docx import Document

    document = Document()
    document.add_heading("采购管理办法", 0)
    document.add_paragraph("第一条 供应商须提交营业执照与审计报告。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "等级"
    table.cell(0, 1).text = "条件"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "两家客户推荐"
    document.save(path)


def _make_pptx(path):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "审批流程"
    left = slide.shapes.add_shape(1, Inches(1), Inches(2), Inches(1), Inches(1))
    left.text_frame.text = "提交申请"
    right = slide.shapes.add_shape(1, Inches(4), Inches(2), Inches(1), Inches(1))
    right.text_frame.text = "分级审批"
    slide.shapes.add_connector(1, Inches(2), Inches(2.5), Inches(4), Inches(2.5))
    deck.save(path)


def _make_xlsx(path):
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "权限矩阵"
    sheet.append(["金额", "审批人"])
    sheet.append(["50 万以下", "部门负责人"])
    book.save(path)


def test_supported_formats_cover_the_common_office_files():
    for suffix in (".pdf", ".docx", ".pptx", ".xlsx", ".md", ".txt", ".html"):
        assert suffix in SUPPORTED_EXT
        assert supported(f"x{suffix}")


def test_format_label_is_human_readable():
    assert format_label("a.docx") == "Word"
    assert format_label("a.pptx") == "PPT"


def test_unsupported_suffix_is_rejected(tmp_path):
    target = tmp_path / "a.zip"
    target.write_bytes(b"PK")
    with pytest.raises(ValueError, match="不支持的格式"):
        extract_pages(target)


def test_missing_file_is_rejected(tmp_path):
    with pytest.raises(ValueError, match="文件不存在"):
        extract_pages(tmp_path / "nope.pdf")


def test_clean_text_normalises_newlines_and_zero_width():
    assert clean_text("a\r\n\r\n\r\nb\u200b") == "a\n\nb"


def test_docx_keeps_tables_as_markdown(workdir):
    target = workdir / "a.docx"
    _make_docx(target)
    pages = extract_pages(target)

    assert pages
    text = "\n".join(page["text"] for page in pages)
    assert "营业执照" in text
    assert "| 等级 | 条件 |" in text


def test_pptx_turns_connectors_into_edges(workdir):
    target = workdir / "a.pptx"
    _make_pptx(target)
    pages = extract_pages(target)

    text = pages[0]["text"]
    assert "提交申请" in text and "分级审批" in text
    assert "提交申请 -> 分级审批" in text


def test_pptx_one_page_per_slide(workdir):
    target = workdir / "b.pptx"
    pytest.importorskip("pptx")
    from pptx import Presentation

    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(target)

    assert len(extract_pages(target)) == 2


def test_xlsx_rows_become_a_markdown_table(workdir):
    target = workdir / "a.xlsx"
    _make_xlsx(target)
    pages = extract_pages(target)

    assert "| 金额 | 审批人 |" in pages[0]["text"]
    assert "部门负责人" in pages[0]["text"]


def test_markdown_and_text_are_split_into_virtual_pages(workdir, monkeypatch):
    from rag.config import settings

    monkeypatch.setattr(settings, "native_page_chars", 40)
    target = workdir / "a.md"
    target.write_text("# 标题\n\n" + "段落内容。" * 20, encoding="utf-8")

    pages = extract_pages(target)
    assert len(pages) > 1
    assert [page["page"] for page in pages] == list(range(1, len(pages) + 1))


def test_html_is_stripped_to_text(workdir):
    target = workdir / "a.html"
    target.write_text(
        "<html><head><style>p{color:red}</style></head>"
        "<body><h1>标题</h1><p>第一段</p><p>第二段</p></body></html>",
        encoding="utf-8",
    )
    text = extract_pages(target)[0]["text"]
    assert "标题" in text and "第一段" in text
    assert "<p>" not in text and "color:red" not in text


def test_csv_becomes_a_table(workdir):
    target = workdir / "a.csv"
    target.write_text("金额,审批人\n50 万,部门负责人\n", encoding="utf-8")
    assert "部门负责人" in extract_pages(target)[0]["text"]
