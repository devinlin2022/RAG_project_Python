"""Service orchestration with every paid call stubbed out."""

import numpy as np
import pytest

import rag.service as service_module
from rag.config import settings
from rag.generate.answer import PageAnswerError, answer_by_pages
from rag.generate.prompt import NO_CONTEXT_ANSWER
from rag.service import Service, ServiceError
from rag.store.db import Database


def _embed_stub(monkeypatch, service):
    """Deterministic vectors: the same text always maps to the same vector."""

    def fake(texts):
        out = []
        for text in texts:
            rng = np.random.default_rng(abs(hash(text)) % (2**32))
            out.append(rng.normal(size=8).tolist())
        return out

    monkeypatch.setattr("rag.service.embed_texts", fake)
    monkeypatch.setattr(service, "_embed_query", lambda query: (fake([query])[0], 1.0, ""))
    return service


def test_ingest_is_idempotent(offline, service, sample_pdf):
    doc, existed = service.ingest(sample_pdf)
    assert existed is False and doc["page_count"] == 3
    again, existed_again = service.ingest(sample_pdf)
    assert existed_again is True and again["id"] == doc["id"]


def test_index_without_embedding_key_still_writes_chunks(offline, service, sample_pdf):
    doc, _ = service.ingest(sample_pdf)
    info = service.index(doc["id"])
    assert info["chunks"] > 0
    assert info["dim"] == 0
    assert service.library.stats()["vectors"] == 0


def test_sparse_retrieval_finds_a_page(offline, service, sample_pdf):
    doc, _ = service.ingest(sample_pdf)
    service.index(doc["id"])
    result = service.retrieve("supplier pricing approval", [doc["id"]], top_k=3)
    assert result["hits"]
    assert result["hits"][0].record.page_from >= 1


def test_dense_request_degrades_to_sparse(offline, service, sample_pdf):
    doc, _ = service.ingest(sample_pdf)
    service.index(doc["id"])
    result = service.retrieve("采购", [doc["id"]], strategy="dense")
    assert "退化为关键词检索" in result["degraded"]
    assert result["strategy"] == "sparse"


def test_ask_returns_sources_and_timings(offline, service, sample_pdf, monkeypatch):
    doc, _ = service.ingest(sample_pdf)
    service.index(doc["id"])
    _embed_stub(monkeypatch, service)
    monkeypatch.setattr(
        "rag.service.chat",
        lambda messages, **kwargs: ("资料显示采购流程有变化 [1]", {"total_tokens": 42}),
    )
    answer = service.ask("what changed in the supplier pricing approval flow?", [doc["id"]], top_k=2)
    assert answer["sources"] and answer["citations"]["valid"] >= 1
    assert answer["timings"]["total_ms"] > 0
    assert "supplier pricing" in service.db.recent_queries(1)[0]["question"]


def test_no_hits_returns_the_abstain_answer(offline, service, sample_pdf, monkeypatch):
    doc, _ = service.ingest(sample_pdf)
    service.index(doc["id"])
    _embed_stub(monkeypatch, service)
    answer = service.ask("supplier pricing", ["missing-doc"], top_k=2)
    assert answer["answer"] == NO_CONTEXT_ANSWER
    assert answer["abstained"] is True
    assert answer["sources"] == []


def test_ask_stream_emits_pipeline_then_answer(offline, service, sample_pdf, monkeypatch):
    doc, _ = service.ingest(sample_pdf)
    service.index(doc["id"])
    _embed_stub(monkeypatch, service)
    monkeypatch.setattr(
        "rag.service.chat_stream", lambda messages, **kwargs: iter(["答", "案 [1]"])
    )

    stages = [event["stage"] for event in service.ask_stream("supplier pricing", [doc["id"]], top_k=2)]
    # The trace the UI renders comes first, then the answer, then the summary.
    assert stages[0] == "embed_start"
    assert "retrieved" in stages
    assert "delta" in stages
    assert stages[-1] == "done"


def test_pipeline_events_carry_measured_numbers(offline, service, sample_pdf, monkeypatch):
    """Every step must report something checkable, or the panel is decoration."""
    doc, _ = service.ingest(sample_pdf)
    service.index(doc["id"])
    _embed_stub(monkeypatch, service)
    monkeypatch.setattr(
        "rag.service.chat_stream", lambda messages, **kwargs: iter(["答 [1]"])
    )

    seen = {}
    for event in service.ask_stream("supplier pricing", [doc["id"]], top_k=2, mode="enhanced"):
        seen[event["stage"]] = event

    assert seen["embed_done"]["dim"] > 0
    assert seen["embed_done"]["norm"] > 0
    assert seen["fuse_done"]["candidates"] >= 1
    assert seen["retrieved"]["sources"]


def test_rewrite_is_used_for_follow_up_questions(offline, service, sample_pdf, monkeypatch):
    doc, _ = service.ingest(sample_pdf)
    service.index(doc["id"])
    _embed_stub(monkeypatch, service)
    monkeypatch.setattr(
        "rag.service.rewrite_question", lambda *a, **k: ("采购流程有哪些变化", True, "")
    )
    monkeypatch.setattr(
        "rag.service.chat", lambda messages, **kwargs: ("结论 [1]", {})
    )
    answer = service.ask(
        "那它呢？", [doc["id"]], top_k=2,
        history=[{"role": "user", "content": "采购流程"}, {"role": "assistant", "content": "…"}],
    )
    assert answer["rewritten"] is True
    assert answer["query"] == "采购流程有哪些变化"


def test_batch_ingest_reports_every_file(offline, service, sample_pdf, tmp_path):
    """A batch must not stop at the first bad file, and order must be preserved."""
    good = tmp_path / "copy.pdf"
    good.write_bytes(sample_pdf.read_bytes())
    broken = tmp_path / "broken.zip"
    broken.write_bytes(b"PK\x03\x04")

    results = service.ingest_files([str(good), str(broken)])

    assert len(results) == 2
    assert results[0]["ok"] is True
    assert results[0]["filename"] == "copy.pdf"
    assert results[1]["ok"] is False
    assert "不支持" in results[1]["error"]


def test_usage_is_accumulated_and_persisted(offline, service, sample_pdf, monkeypatch):
    doc, _ = service.ingest(sample_pdf)
    service.index(doc["id"])
    _embed_stub(monkeypatch, service)
    monkeypatch.setattr(
        "rag.service.chat",
        lambda messages, **kwargs: ("结论 [1]", {"prompt_tokens": 100, "completion_tokens": 20}),
    )

    answer = service.ask("supplier pricing", [doc["id"]], top_k=2)

    assert answer["usage"]["total_tokens"] == 120
    assert service.usage_summary()["llm_in"] >= 100
    # A restart must not reset the counter: it is summed from the query log.
    assert service.db.usage_totals()["llm_out"] >= 20


def test_two_modes_differ_in_cost_and_depth(offline, service):
    fast = service.resolve_mode("fast")
    enhanced = service.resolve_mode("enhanced")

    assert fast["rerank"] is False and fast["rewrite"] is False
    assert enhanced["rerank"] is True and enhanced["rewrite"] is True
    assert enhanced["candidates"] > fast["candidates"]


def test_deck_is_chunked_per_slide(offline, service, tmp_path):
    """A slide deck must never collapse into one block citing p.1-3."""
    pytest.importorskip("pptx")
    from pptx import Presentation

    deck = Presentation()
    for title in ("采购申请", "预算校验", "分级审批"):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text_frame.text = f"{title}环节说明与责任部门。"
    path = tmp_path / "deck.pptx"
    deck.save(path)

    doc, _ = service.ingest(path)
    info = service.index(doc["id"])
    chunks = [c for c in service.library.records if c.doc_id == doc["id"]]

    assert info["strategy"] == "per-page"
    assert len({chunk.page_from for chunk in chunks}) == 3
    assert all(chunk.page_from == chunk.page_to for chunk in chunks)


def test_delete_removes_document_and_chunks(offline, service, sample_pdf):
    doc, _ = service.ingest(sample_pdf)
    service.index(doc["id"])
    assert service.delete(doc["id"]) is True
    assert service.library.stats()["chunks"] == 0
    with pytest.raises(ServiceError):
        service.get(doc["id"])


def test_documents_survive_a_restart(offline, sample_pdf, tmp_path):
    db_path = tmp_path / "restart.sqlite"
    first = Service(Database(db_path))
    doc, _ = first.ingest(sample_pdf)
    first.index(doc["id"])
    chunks = first.library.stats()["chunks"]

    second = Service(Database(db_path))
    assert doc["id"] in second.docs
    assert second.library.stats()["chunks"] == chunks
    assert second.docs[doc["id"]]["page_count"] == 3


def test_page_qa_rejects_empty_page_selection(offline, service, sample_pdf):
    doc, _ = service.ingest(sample_pdf)
    with pytest.raises(PageAnswerError):
        answer_by_pages(doc, [], "问题")


def test_unknown_strategy_is_rejected(offline, service, sample_pdf):
    doc, _ = service.ingest(sample_pdf)
    service.index(doc["id"])
    with pytest.raises(ServiceError):
        service.retrieve("x", [doc["id"]], strategy="magic")


def test_rerank_path_is_wired(offline, service, sample_pdf, monkeypatch):
    """A reranker is only called when a key exists -- and its order must win."""
    doc, _ = service.ingest(sample_pdf)
    service.index(doc["id"])
    _embed_stub(monkeypatch, service)
    monkeypatch.setattr("rag.service.settings.llm_api_key", "sk-test")
    monkeypatch.setattr("rag.service.settings.rerank_enabled", True)
    monkeypatch.setattr("rag.service.reranker.rerank", lambda query, docs: [
        (index, float(index)) for index in range(len(docs))
    ])
    result = service.retrieve("supplier pricing", [doc["id"]], top_k=2, rerank=True)
    assert result["rerank"]["used"] is True
    assert result["rerank"]["model"] == settings.rerank_model
    assert result["hits"][0].rerank is not None
